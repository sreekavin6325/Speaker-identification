from __future__ import annotations

import csv
import json
import math
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PAPER_DIR = ROOT / "paper"
OUT_DOCX = PAPER_DIR / "Multi_Model_Speaker_Recognition_IEEE_Paper.docx"
PROTOCOL = "librispeech_dev_clean_chapter_heldout_2fold__01ff6bcaaebb"

BASELINE = ROOT / "research_results" / "baseline_cv" / "protocols" / PROTOCOL / "summary"
ROBUST = (
    ROOT
    / "research_results"
    / "robustness_benchmark"
    / "protocols"
    / PROTOCOL
    / "summary"
    / "logistic_regression"
    / "full"
)
PREPROC = (
    ROOT
    / "research_results"
    / "preprocessing_benchmark"
    / "protocols"
    / PROTOCOL
    / "summary"
    / "logistic_regression"
    / "full"
)
TUNING = (
    ROOT
    / "research_results"
    / "hyperparameter_tuning"
    / "protocols"
    / PROTOCOL
    / "summary"
    / "svm"
    / "full"
)


FIG_CLASSIFIER = BASELINE / "classifier_benchmark" / "full" / "accuracy_heatmap.png"
FIG_SHORT = ROBUST / "short_duration_accuracy.png"
FIG_NOISE = ROBUST / "noise_accuracy.png"
FIG_TIMING = BASELINE / "linear_svm" / "full" / "warm_extraction_timing_bar.png"
FIG_PREPROC = PREPROC / "preprocessing_accuracy_heatmap.png"
FIG_EFFECTS = PREPROC / "preprocessing_main_effects.png"


BLUE = "1F4E79"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "E7E6E6"
MID_GRAY = "BFBFBF"
WHITE = "FFFFFF"
BLACK = "000000"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "A6A6A6", size: int = 4) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=45, start=55, bottom=45, end=55) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text: str, *, bold=False, color=BLACK, size=7.2, align=None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def set_column_width(cell, width_inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float] | None = None,
    font_size: float = 7.2,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, BLUE)
        set_cell_text(cell, header, bold=True, color=WHITE, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_border(cell, color=WHITE, size=5)
        set_cell_margins(cell)
        if widths:
            set_column_width(cell, widths[idx])
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        for col_idx, value in enumerate(values):
            cell = row.cells[col_idx]
            set_cell_shading(cell, WHITE if row_idx % 2 == 0 else "F4F7FA")
            set_cell_text(cell, value, size=font_size)
            set_cell_border(cell, color="C8C8C8", size=3)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                set_column_width(cell, widths[col_idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def set_section_layout(section, columns: int) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.3)
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(columns))
    cols.set(qn("w:space"), "300")


def switch_columns(doc: Document, columns: int) -> None:
    section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_section_layout(section, columns)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run.font.name = "Times New Roman"
    run.font.size = Pt(8)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(6 if level == 1 else 4)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text.upper() if level == 1 else text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11.5 if level == 1 else 10.5)
    run.bold = True
    if level == 1:
        run.font.small_caps = True


def add_body(doc: Document, text: str, *, first_line=True) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.space_after = Pt(2)
    if first_line:
        paragraph.paragraph_format.first_line_indent = Inches(0.16)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.14)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(9.8)


def add_equation(doc: Document, equation: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(equation)
    run.italic = True
    run.font.name = "Cambria Math"
    run.font.size = Pt(10)


def add_caption(doc: Document, text: str, *, table=False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = table
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(8)


def add_wide_figure(doc: Document, path: Path, caption: str, width: float = 7.05) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_native_pipeline(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=11)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    boxes = [
        "LibriSpeech\nmanifest",
        "→",
        "Chapter-held-out\n2-fold protocol",
        "→",
        "Four frozen\nembedding models",
        "→",
        "Eight\nclassifiers",
        "→",
        "Robustness &\npreprocessing",
        "→",
        "Accuracy, F1,\nlatency",
    ]
    widths = [0.86, 0.23, 1.10, 0.23, 1.05, 0.23, 0.78, 0.23, 1.03, 0.23, 0.90]
    for i, (cell, text) in enumerate(zip(table.rows[0].cells, boxes)):
        set_column_width(cell, widths[i])
        set_cell_margins(cell, top=75, start=35, bottom=75, end=35)
        if text == "→":
            set_cell_shading(cell, WHITE)
            set_cell_text(cell, text, bold=True, color=BLUE, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_border(cell, color=WHITE, size=0)
        else:
            set_cell_shading(cell, LIGHT_BLUE)
            set_cell_text(cell, text, bold=True, color=BLUE, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_border(cell, color=BLUE, size=5)
    add_caption(doc, "Fig. 1. Experimental workflow and leakage-controlled evaluation pipeline.")


def configure_document(doc: Document) -> None:
    set_section_layout(doc.sections[0], 1)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0
    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(9)


def add_title_block(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    run = p.add_run(
        "Comparative Evaluation of Pre-Trained Speaker Embeddings for Closed-Set "
        "Speaker Identification Under Classifier, Robustness, Efficiency, and "
        "Preprocessing Variations"
    )
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)
    run.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("Sree Kavin")
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    run = p.add_run("Independent Research Project · Affiliation and email to be supplied for submission")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    run.italic = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.right_indent = Inches(0.35)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Abstract—")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(9)
    r = p.add_run(
        "This work presents a controlled comparison of four pre-trained speaker embedding "
        "models—ECAPA-TDNN, X-Vector, WavLM Base Plus for speaker verification, and "
        "UniSpeech-SAT Base Plus for speaker verification—for closed-set speaker identification. "
        "The study evaluates eight downstream classifiers, nested support-vector-machine tuning, "
        "short-duration speech, additive white noise, CPU embedding latency, and a complete "
        "2³ preprocessing factorial. A strict chapter-held-out two-fold protocol was constructed "
        "from LibriSpeech dev-clean, comprising 2,185 clips from 31 speakers and 88 source chapters. "
        "The best single system was ECAPA-TDNN with cosine-centroid classification at 99.00% mean "
        "accuracy. Logistic regression was the most consistently effective classifier, ranking first "
        "for the other three embeddings and averaging 95.52% across models. ECAPA-TDNN was also "
        "the most robust model at every tested duration and signal-to-noise ratio, retaining 82.85% "
        "accuracy at 0.5 s and 89.17% at 0 dB. X-Vector offered the best CPU speed/accuracy trade-off "
        "with a 30.54 ms warm median extraction time and 95.65% linear-SVM accuracy. Preprocessing "
        "effects were model dependent rather than universal. The findings demonstrate that model "
        "selection must jointly consider recognition accuracy, robustness, latency, and front-end "
        "processing rather than relying on a single clean-condition score."
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(9)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.right_indent = Inches(0.35)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("Index Terms—")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(9)
    r = p.add_run(
        "speaker identification, speaker embeddings, ECAPA-TDNN, X-Vector, WavLM, "
        "UniSpeech-SAT, classifier benchmarking, robustness, preprocessing ablation"
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(9)


def build_paper() -> None:
    PAPER_DIR.mkdir(parents=True, exist_ok=True)

    best_rows = read_csv(BASELINE / "classifier_benchmark" / "full" / "best_classifier_per_model.csv")
    efficiency_rows = read_csv(BASELINE / "linear_svm" / "full" / "model_comparison.csv")
    robust_rows = read_csv(ROBUST / "model_condition_summary.csv")
    preproc_rows = read_csv(PREPROC / "model_variant_summary.csv")
    effects_rows = read_csv(PREPROC / "factorial_effects_summary.csv")
    tuning_rows = read_csv(TUNING / "model_level_comparison.csv")

    doc = Document()
    configure_document(doc)
    add_title_block(doc)

    add_heading(doc, "I. Introduction")
    add_body(
        doc,
        "Speaker identification assigns a speech segment to one of a set of enrolled speakers. "
        "Modern systems usually separate representation learning from the final identity decision: "
        "a pre-trained neural network converts speech into a fixed-dimensional speaker embedding, "
        "and a comparatively lightweight classifier maps the embedding to an enrolled identity. "
        "This modular design makes it possible to reuse a common embedding model while comparing "
        "classifiers, deployment costs, and robustness without re-training the neural backbone."
    )
    add_body(
        doc,
        "Published accuracy values are often difficult to compare because they mix datasets, split "
        "rules, preprocessing, classifiers, and hardware. Random clip-level splitting is particularly "
        "risky for audiobooks: clips originating from the same chapter can share channel, room, and "
        "recording characteristics. A model may therefore exploit source similarity instead of true "
        "speaker identity. This study addresses that threat by holding out source chapters and applying "
        "the same folds to every model and classifier."
    )
    add_body(
        doc,
        "The goal is not merely to name a winner. Five research questions are investigated: (RQ1) "
        "which pre-trained embedding produces the highest recognition accuracy; (RQ2) which classifier "
        "is most effective for speaker embeddings; (RQ3) which model is most robust to short speech and "
        "noise; (RQ4) which model best balances recognition and computational efficiency; and (RQ5) "
        "which preprocessing operations most affect performance."
    )
    add_body(doc, "The principal contributions are:", first_line=False)
    add_bullet(doc, "a shared, leakage-controlled, chapter-held-out evaluation protocol for all experimental conditions;")
    add_bullet(doc, "a four-embedding by eight-classifier benchmark, supplemented by nested SVM tuning;")
    add_bullet(doc, "controlled short-duration, additive-noise, CPU-latency, and 2³ preprocessing experiments; and")
    add_bullet(doc, "an explicit deployment-oriented interpretation of accuracy, robustness, and efficiency trade-offs.")

    add_heading(doc, "II. Related Work")
    add_body(
        doc,
        "X-Vectors use a time-delay neural network and statistics pooling to obtain utterance-level "
        "representations [1]. ECAPA-TDNN extends the TDNN family with channel attention, hierarchical "
        "feature aggregation, and Res2Net-style modules [2]. Both evaluated SpeechBrain checkpoints were "
        "pre-trained for speaker recognition on VoxCeleb-style data and used here as frozen feature "
        "extractors through the SpeechBrain toolkit [3], [4]."
    )
    add_body(
        doc,
        "Self-supervised speech encoders learn general acoustic representations from large unlabeled "
        "corpora. WavLM introduces masked speech prediction with denoising and has demonstrated strong "
        "SUPERB performance [5], [6]. UniSpeech-SAT explicitly incorporates speaker-aware pre-training "
        "and utterance-level contrastive objectives [7]. Their speaker-verification checkpoints include "
        "X-Vector-style heads and yield 512-dimensional embeddings in this work."
    )
    add_body(
        doc,
        "Unlike studies focused on a single architecture, the present experiment holds the data protocol "
        "constant while varying both the embedding and the downstream classifier. It further treats "
        "robustness and preprocessing as first-class evaluation dimensions."
    )

    add_heading(doc, "III. Experimental Methodology")
    add_heading(doc, "A. Dataset and Strict Evaluation Protocol", level=2)
    add_body(
        doc,
        "LibriSpeech dev-clean [8] was converted into a research manifest. Speakers with insufficient "
        "chapter diversity for the strict protocol were excluded. The resulting set contains 2,185 clips, "
        "31 speakers, 88 source chapters, and 4.174 h of speech. Clip duration has a mean of 6.877 s, a "
        "median of 5.780 s, and a range of 1.535–32.485 s. Speaker clip counts range from 38 to 101."
    )
    add_body(
        doc,
        "Two deterministic outer folds were constructed at the chapter level. Fold test sizes were 1,069 "
        "and 1,116 clips. A clip or source chapter assigned to an outer test partition never appeared in "
        "its corresponding training partition. Because the same corpus is reused across two correlated "
        "folds, fold variability is reported descriptively and no formal significance claim is made."
    )
    add_caption(doc, "TABLE I\nSTRICT DATASET AND EVALUATION PROTOCOL", table=True)
    add_table(
        doc,
        ["Property", "Value"],
        [
            ["Task", "Closed-set speaker identification"],
            ["Corpus", "LibriSpeech dev-clean"],
            ["Clips / speakers / chapters", "2,185 / 31 / 88"],
            ["Total duration", "4.174 h"],
            ["Clip duration", "Mean 6.877 s; median 5.780 s; 1.535–32.485 s"],
            ["Outer protocol", "Deterministic chapter-held-out 2-fold"],
            ["Fold test sizes", "1,069 and 1,116 clips"],
            ["Leakage checks", "Zero shared clip paths and zero shared chapters"],
        ],
        widths=[1.20, 2.10],
        font_size=7.3,
    )
    add_native_pipeline(doc)

    add_heading(doc, "B. Pre-Trained Speaker Embeddings", level=2)
    add_body(
        doc,
        "All neural backbones were frozen. No ECAPA-TDNN, X-Vector, WavLM, or UniSpeech-SAT network "
        "weights were updated on the evaluation corpus; only the downstream closed-set classifiers were "
        "fitted. Audio was converted to mono and resampled to 16 kHz before embedding extraction. "
        "Embeddings were L2-normalized where required by the evaluation pipeline."
    )
    add_caption(doc, "TABLE II\nEVALUATED PRE-TRAINED EMBEDDING MODELS", table=True)
    add_table(
        doc,
        ["Model", "Architecture", "Dim.", "Pre-trained checkpoint"],
        [
            ["ECAPA-TDNN", "Attentive TDNN", "192", "speechbrain/spkrec-ecapa-voxceleb"],
            ["X-Vector", "TDNN + statistics pooling", "512", "speechbrain/spkrec-xvect-voxceleb"],
            ["WavLM", "Transformer + X-Vector head", "512", "microsoft/wavlm-base-plus-sv"],
            ["UniSpeech-SAT", "SSL encoder + X-Vector head", "512", "microsoft/unispeech-sat-base-plus-sv"],
        ],
        widths=[0.75, 1.05, 0.38, 1.35],
        font_size=6.8,
    )

    add_heading(doc, "C. Classifiers and Metrics", level=2)
    add_body(
        doc,
        "Eight downstream classifiers were compared: linear SVM, radial-basis-function SVM, logistic "
        "regression, k-nearest neighbors, random forest, decision tree, XGBoost, and cosine centroid. "
        "Feature scaling was learned only from each outer-training partition. Hyperparameters in the "
        "separate SVM-tuning study were selected within the outer-training data; untouched outer-test "
        "chapters remained the final evaluation data."
    )
    add_equation(doc, "Accuracy = N(correct predictions) / N(test predictions)")
    add_body(
        doc,
        "Mean outer-fold accuracy is the primary measure, with macro-F1 used to verify performance across "
        "speakers. Classifier fit time, prediction time, and embedding extraction latency were measured "
        "separately. Cold-start time includes the first successful inference; warm latency summarizes "
        "subsequent CPU extraction."
    )

    add_heading(doc, "D. Robustness and Preprocessing Experiments", level=2)
    add_body(
        doc,
        "For robustness, logistic-regression classifiers were trained on clean embeddings and evaluated on "
        "held-out clips truncated to 0.5, 1, 2, and 3 s. A separate experiment added deterministic white "
        "noise at 20, 10, and 0 dB SNR to held-out speech. Noise scaling followed the signal/noise power "
        "ratio, and transformed test clips never entered training."
    )
    add_equation(doc, "SNR(dB) = 10 log10(Psignal / Pnoise)")
    add_body(
        doc,
        "The preprocessing study used a complete 2³ factorial over silence trimming, stationary spectral "
        "noise reduction, and peak normalization. The fixed processing order was trim → denoise → "
        "normalize. Trimming used a frame-RMS threshold 30 dB below the maximum, denoising used a "
        "stationary gate (proportion decrease 0.8), and normalization set peak magnitude to 0.95."
    )

    add_heading(doc, "IV. Results")
    add_heading(doc, "A. Embedding and Classifier Comparison", level=2)
    add_body(
        doc,
        "Table III reports the classifier selected independently for each embedding. ECAPA-TDNN with "
        "cosine centroid achieved the highest overall mean accuracy (99.00%) and macro-F1 (98.98%). "
        "Logistic regression ranked first for X-Vector, WavLM, and UniSpeech-SAT and had the highest "
        "average accuracy across the four embeddings (95.52%). Linear SVM ranked second by cross-model "
        "average (94.68%). The decision tree was consistently weakest, averaging 63.15%."
    )
    add_caption(doc, "TABLE III\nBEST CLASSIFIER FOR EACH EMBEDDING MODEL", table=True)
    table_rows = []
    for row in best_rows:
        table_rows.append(
            [
                row["model_display_name"],
                row["best_classifier_display_name"],
                f"{float(row['accuracy_mean_percent']):.2f}",
                f"{float(row['f1_macro_mean_percent']):.2f}",
                f"{float(row['classifier_fit_seconds_mean']):.3f}",
            ]
        )
    add_table(
        doc,
        ["Embedding", "Best classifier", "Acc. (%)", "Macro-F1 (%)", "Fit (s)"],
        table_rows,
        widths=[0.78, 0.90, 0.52, 0.62, 0.45],
        font_size=6.7,
    )
    add_wide_figure(
        doc,
        FIG_CLASSIFIER,
        "Fig. 2. Mean chapter-held-out recognition accuracy for four embeddings and eight classifiers.",
        width=7.05,
    )
    add_body(
        doc,
        "Nested SVM tuning did not produce a universal gain. Relative to the fixed linear SVM, tuned "
        "accuracy changed by 0.00, +0.05, −0.85, and +0.09 percentage points for ECAPA-TDNN, "
        "X-Vector, WavLM, and UniSpeech-SAT, respectively. This result reinforces the need to evaluate "
        "simple linear baselines and to avoid assuming that a larger search automatically generalizes."
    )

    add_heading(doc, "B. Short-Duration and Noise Robustness", level=2)
    add_body(
        doc,
        "ECAPA-TDNN led at every tested duration. At 0.5 s it retained 82.85% accuracy, compared with "
        "63.09% for X-Vector, 29.02% for WavLM, and 24.28% for UniSpeech-SAT. At 1 s, ECAPA-TDNN "
        "reached 96.66%; both self-supervised systems required longer speech to approach their clean "
        "condition. ECAPA-TDNN reached 98.91% at 2 s and remained essentially saturated at 3 s."
    )
    add_wide_figure(
        doc,
        FIG_SHORT,
        "Fig. 3. Mean outer-fold accuracy under controlled test-segment truncation.",
        width=6.95,
    )
    add_body(
        doc,
        "The noise experiment produced an even larger separation. At 20 and 10 dB SNR, ECAPA-TDNN "
        "retained 98.21% and 97.42%, respectively, and at 0 dB retained 89.17%. X-Vector fell to "
        "14.84% at 0 dB, while WavLM and UniSpeech-SAT reached 22.98% and 26.53%. The latter ranking "
        "suggests that the self-supervised systems resisted severe noise slightly better than X-Vector, "
        "but none approached ECAPA-TDNN."
    )
    add_caption(doc, "TABLE IV\nROBUSTNESS ACCURACY (%)", table=True)
    robust_lookup: dict[str, dict[str, float]] = {}
    for row in robust_rows:
        model = row.get("model_display_name", row.get("model_id", ""))
        condition = row.get("condition_id", row.get("condition", ""))
        value = row.get("accuracy_mean_percent")
        if value is None or value == "":
            value = str(float(row["accuracy_mean"]) * 100)
        robust_lookup.setdefault(model, {})[condition] = float(value)
    preferred_names = [
        ("ECAPA-TDNN", "speechbrain_ecapa"),
        ("X-Vector", "speechbrain_xvector"),
        ("WavLM", "wavlm_base_plus_sv"),
        ("UniSpeech-SAT", "unispeech_sat_base_plus_sv"),
    ]
    condition_aliases = {
        "Clean": ["clean", "raw"],
        "0.5 s": ["short_0p5s", "0.5s"],
        "1 s": ["short_1s", "1s"],
        "2 s": ["short_2s", "2s"],
        "3 s": ["short_3s", "3s"],
        "20 dB": ["noise_white_20db", "20db"],
        "10 dB": ["noise_white_10db", "10db"],
        "0 dB": ["noise_white_0db", "0db"],
    }
    def find_metric(display: str, model_id: str, aliases: list[str]) -> float:
        for key, values in robust_lookup.items():
            if display.lower() in key.lower() or model_id.lower() in key.lower():
                for alias in aliases:
                    if alias in values:
                        return values[alias]
                for condition, value in values.items():
                    if any(alias.lower() in condition.lower() for alias in aliases):
                        return value
        return math.nan
    robust_table = []
    for display, model_id in preferred_names:
        values = [find_metric(display, model_id, aliases) for aliases in condition_aliases.values()]
        robust_table.append([display] + ["—" if math.isnan(v) else f"{v:.2f}" for v in values])
    add_table(
        doc,
        ["Model"] + list(condition_aliases.keys()),
        robust_table,
        widths=[0.70, 0.38, 0.38, 0.38, 0.38, 0.38, 0.42, 0.42, 0.42],
        font_size=5.9,
    )
    add_wide_figure(
        doc,
        FIG_NOISE,
        "Fig. 4. Mean outer-fold accuracy under additive white noise. Lower SNR is more difficult.",
        width=6.95,
    )

    add_heading(doc, "C. Computational Efficiency", level=2)
    add_body(
        doc,
        "Efficiency was measured on CPU under the same extraction pipeline. X-Vector had the lowest warm "
        "median latency (30.54 ms), followed by ECAPA-TDNN (134.43 ms), UniSpeech-SAT (387.44 ms), "
        "and WavLM (444.18 ms). Under a shared linear-SVM classifier, ECAPA-TDNN produced the highest "
        "accuracy (98.40%), while X-Vector achieved 95.65% at approximately one quarter of ECAPA's warm "
        "median extraction time. ECAPA-TDNN and X-Vector therefore form the practical accuracy/latency "
        "frontier in this CPU experiment."
    )
    add_caption(doc, "TABLE V\nCPU EFFICIENCY UNDER A SHARED LINEAR SVM", table=True)
    eff_table = []
    for row in efficiency_rows:
        eff_table.append(
            [
                row["model_display_name"].replace("SpeechBrain — ", "").replace("Microsoft — ", ""),
                row["embedding_dimension"],
                f"{float(row['accuracy_mean_percent']):.2f}",
                f"{float(row['warm_extraction_median_ms']):.2f}",
                f"{float(row['warm_extraction_p95_ms']):.2f}",
                f"{float(row['cold_first_success_ms']):.1f}",
            ]
        )
    add_table(
        doc,
        ["Model", "Dim.", "Acc. (%)", "Warm med. (ms)", "Warm P95 (ms)", "Cold (ms)"],
        eff_table,
        widths=[0.76, 0.36, 0.48, 0.66, 0.66, 0.58],
        font_size=6.2,
    )
    add_wide_figure(
        doc,
        FIG_TIMING,
        "Fig. 5. Warm CPU embedding-extraction latency. Bars show mean latency; labels report the median.",
        width=6.85,
    )

    add_heading(doc, "D. Preprocessing Ablation", level=2)
    add_body(
        doc,
        "No preprocessing configuration was universally optimal. Silence trimming was beneficial for "
        "X-Vector, improving logistic-regression accuracy from 96.40% to 97.43% (+1.03 percentage "
        "points), but reduced the average factorial response for WavLM and UniSpeech-SAT. WavLM's best "
        "variant was denoise plus normalization at 94.40% (+1.20 points over raw). UniSpeech-SAT's best "
        "variant was the raw input at 93.57%. ECAPA-TDNN changed only marginally, peaking at 98.96% "
        "with trimming."
    )
    best_preproc = []
    grouped: dict[str, list[dict[str, str]]] = {}
    for row in preproc_rows:
        grouped.setdefault(row["model_id"], []).append(row)
    display_map = {
        "speechbrain_ecapa": "ECAPA-TDNN",
        "speechbrain_xvector": "X-Vector",
        "wavlm_base_plus_sv": "WavLM",
        "unispeech_sat_base_plus_sv": "UniSpeech-SAT",
    }
    for model_id in display_map:
        group = grouped[model_id]
        best = max(group, key=lambda item: float(item["accuracy_mean"]))
        raw = next(item for item in group if item["variant"] == "raw")
        delta = 100 * (float(best["accuracy_mean"]) - float(raw["accuracy_mean"]))
        best_preproc.append(
            [
                display_map[model_id],
                best["variant"].replace("_", " + "),
                f"{100 * float(best['accuracy_mean']):.2f}",
                f"{delta:+.2f}",
            ]
        )
    add_caption(doc, "TABLE VI\nBEST PREPROCESSING VARIANT PER MODEL", table=True)
    add_table(
        doc,
        ["Model", "Best variant", "Acc. (%)", "Δ vs. raw (pp)"],
        best_preproc,
        widths=[0.80, 1.15, 0.58, 0.70],
        font_size=6.8,
    )
    add_wide_figure(
        doc,
        FIG_PREPROC,
        "Fig. 6. Complete 2³ preprocessing ablation under the strict protocol.",
        width=7.0,
    )
    add_wide_figure(
        doc,
        FIG_EFFECTS,
        "Fig. 7. Factorial main effects of trimming, denoising, and peak normalization. Positive values indicate improvement.",
        width=6.95,
    )

    add_heading(doc, "V. Discussion")
    add_heading(doc, "A. Answers to the Research Questions", level=2)
    add_body(
        doc,
        "RQ1—Highest recognition accuracy: ECAPA-TDNN was the strongest embedding. Its cosine-centroid "
        "system achieved 99.00%, and it also ranked first under every common classifier except the "
        "comparison is not meaningful where classifier families differ in inductive bias. The result "
        "supports ECAPA-TDNN as the accuracy-first choice for this clean English audiobook domain."
    )
    add_body(
        doc,
        "RQ2—Most effective classifier: Logistic regression was the most reliable general classifier, "
        "ranking first for three of four embeddings and averaging 95.52%. Cosine centroid achieved the "
        "single best result only with ECAPA-TDNN. Consequently, logistic regression is the strongest "
        "default when one classifier must be shared, while cosine scoring should be considered for "
        "well-separated, angularly structured ECAPA embeddings."
    )
    add_body(
        doc,
        "RQ3—Robustness: ECAPA-TDNN was best under every short-duration and noise condition. Its 0.5-s "
        "and 0-dB results were 19.76 and 62.64 percentage points above the next-best systems, "
        "respectively. The advantage is large enough to affect model selection for voice interfaces, "
        "brief commands, and noisy capture."
    )
    add_body(
        doc,
        "RQ4—Performance/efficiency balance: X-Vector is the most attractive CPU-constrained option. It "
        "sacrificed 2.75 points of linear-SVM accuracy relative to ECAPA-TDNN but reduced warm median "
        "extraction latency by 77.3%. ECAPA-TDNN remains preferable when recognition and robustness are "
        "more important than latency."
    )
    add_body(
        doc,
        "RQ5—Preprocessing impact: Effects were architecture dependent. X-Vector benefited most from "
        "silence trimming; WavLM benefited most from denoising plus normalization; UniSpeech-SAT was "
        "best without added processing; ECAPA-TDNN was insensitive within about one tenth of a point. "
        "A single global preprocessing recipe is therefore not scientifically justified by these data."
    )

    add_heading(doc, "B. Practical System Selection", level=2)
    add_body(
        doc,
        "For a high-accuracy desktop application, ECAPA-TDNN with either cosine centroid or logistic "
        "regression is the recommended primary model. For faster CPU inference, X-Vector with logistic "
        "regression and silence trimming provides a better latency/accuracy compromise. WavLM and "
        "UniSpeech-SAT remain useful research comparators and may respond differently under broader "
        "domain adaptation, multilingual speech, or task-specific fine-tuning, but they were dominated "
        "by the two TDNN-family systems in the present frozen-embedding evaluation."
    )

    add_heading(doc, "VI. Limitations and Threats to Validity")
    add_body(
        doc,
        "The study is closed-set: every test speaker is enrolled during classifier training. A high "
        "closed-set class probability is not a calibrated guarantee that an unknown speaker belongs to "
        "the enrolled set. Open-set deployment requires an explicit rejection score and threshold, tuned "
        "on disjoint known and unknown validation speakers, followed by reporting of false-accept rate, "
        "false-reject rate, equal-error rate, and detection cost."
    )
    add_body(
        doc,
        "The evaluation uses one clean English audiobook subset. Its two folds share the same speaker "
        "population and are correlated; confidence intervals based on only two folds are not reliable for "
        "formal inference. White noise is a controlled stressor rather than a substitute for real rooms, "
        "microphones, reverberation, music, or competing speech. CPU latency depends on the exact Windows "
        "computer, software stack, and caching state. Finally, the neural encoders were frozen, so the "
        "results do not measure the potential benefits or costs of task-specific fine-tuning."
    )

    add_heading(doc, "VII. Reproducibility")
    add_body(
        doc,
        f"Every experiment used protocol identifier {PROTOCOL} and protocol SHA-256 prefix "
        "01ff6bcaaebb. Manifests, embedding metadata, fold assignments, per-condition results, and "
        "summary CSV/JSON artifacts were retained. All robustness and preprocessing embeddings achieved "
        "100% coverage. The preprocessing factorial required 61,180 embedding extractions. Deterministic "
        "fold assignments and shared condition definitions enable direct model comparisons."
    )

    add_heading(doc, "VIII. Conclusion")
    add_body(
        doc,
        "This work compared four pre-trained speaker embeddings under a common leakage-controlled protocol "
        "and evaluated classifier choice, SVM tuning, short speech, additive noise, CPU cost, and "
        "preprocessing. ECAPA-TDNN was the overall accuracy and robustness leader; cosine centroid "
        "produced the best single system at 99.00%, while logistic regression was the most consistent "
        "cross-model classifier. X-Vector delivered the best CPU-oriented speed/accuracy compromise. "
        "Preprocessing gains were model specific and nested SVM tuning did not reliably improve results. "
        "Future work should add open-set rejection, independent evaluation speakers and corpora, real "
        "environmental noise and reverberation, calibrated uncertainty, multilingual speech, and enough "
        "independent folds or repeated trials for statistical significance testing."
    )

    add_heading(doc, "References")
    references = [
        "[1] D. Snyder, D. Garcia-Romero, G. Sell, D. Povey, and S. Khudanpur, “X-vectors: Robust DNN embeddings for speaker recognition,” in Proc. ICASSP, 2018, pp. 5329–5333, doi: 10.1109/ICASSP.2018.8461375.",
        "[2] B. Desplanques, J. Thienpondt, and K. Demuynck, “ECAPA-TDNN: Emphasized channel attention, propagation and aggregation in TDNN based speaker verification,” in Proc. Interspeech, 2020, pp. 3830–3834, doi: 10.21437/Interspeech.2020-2650.",
        "[3] M. Ravanelli et al., “SpeechBrain: A general-purpose speech toolkit,” arXiv:2106.04624, 2021.",
        "[4] A. Nagrani, J. S. Chung, and A. Zisserman, “VoxCeleb: A large-scale speaker identification dataset,” in Proc. Interspeech, 2017, pp. 2616–2620, doi: 10.21437/Interspeech.2017-950.",
        "[5] S. Chen et al., “WavLM: Large-scale self-supervised pre-training for full stack speech processing,” IEEE J. Sel. Topics Signal Process., vol. 16, no. 6, pp. 1505–1518, 2022, doi: 10.1109/JSTSP.2022.3188113.",
        "[6] S.-w. Yang et al., “SUPERB: Speech Processing Universal PERformance Benchmark,” in Proc. Interspeech, 2021, pp. 1194–1198.",
        "[7] S. Chen et al., “UniSpeech-SAT: Universal speech representation learning with speaker aware pre-training,” in Proc. ICASSP, 2022.",
        "[8] V. Panayotov, G. Chen, D. Povey, and S. Khudanpur, “LibriSpeech: An ASR corpus based on public domain audio books,” in Proc. ICASSP, 2015, pp. 5206–5210, doi: 10.1109/ICASSP.2015.7178964.",
        "[9] C. Cortes and V. Vapnik, “Support-vector networks,” Mach. Learn., vol. 20, pp. 273–297, 1995.",
        "[10] F. Pedregosa et al., “Scikit-learn: Machine learning in Python,” J. Mach. Learn. Res., vol. 12, pp. 2825–2830, 2011.",
        "[11] N. Dehak, P. J. Kenny, R. Dehak, P. Dumouchel, and P. Ouellet, “Front-end factor analysis for speaker verification,” IEEE Trans. Audio, Speech, Lang. Process., vol. 19, no. 4, pp. 788–798, 2011.",
        "[12] J. H. L. Hansen and T. Hasan, “Speaker recognition by machines and humans: A tutorial review,” IEEE Signal Process. Mag., vol. 32, no. 6, pp. 74–99, 2015.",
    ]
    for reference in references:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(reference)
        run.font.name = "Times New Roman"
        run.font.size = Pt(7.6)

    doc.core_properties.title = (
        "Comparative Evaluation of Pre-Trained Speaker Embeddings for Closed-Set Speaker Identification"
    )
    doc.core_properties.author = "Sree Kavin"
    doc.core_properties.subject = "IEEE-style research manuscript"
    doc.core_properties.keywords = (
        "speaker identification, ECAPA-TDNN, X-Vector, WavLM, UniSpeech-SAT, robustness"
    )

    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build_paper()
